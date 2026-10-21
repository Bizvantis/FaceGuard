"""Build an editable, IEEE-formatted Word version of faceguard.tex.

    uv run --no-project --with python-docx --with pymupdf python make_docx.py BUILD_DIR OUT.docx

BUILD_DIR must contain faceguard.aux and faceguard.bbl from a LaTeX run of this folder
(e.g. `tectonic --keep-intermediates faceguard.tex`), so that section, figure, table and citation
numbers match the PDF exactly. Needs pandoc on PATH and a TeX engine (TECTONIC env var or `tectonic`)
for the two TikZ diagrams.

Pipeline: expand macros / resolve refs and cites -> pandoc (LaTeX -> docx, native Word equations)
-> python-docx post-processing (IEEE page layout, styles, numbered headings, captions, table rules,
numbered equations, figures, references, running header).
"""

from __future__ import annotations

import copy
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import docx
import pymupdf
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips

HERE = Path(__file__).resolve().parent
ROMAN = ["", "I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII", "XIII", "XIV", "XV"]

# IEEE journal geometry (US Letter): two 3.5-in columns, 0.16-in gutter
PAGE_W, PAGE_H = 12240, 15840
MARGIN_LR, MARGIN_T, MARGIN_B = 965, 1080, 1080
GUTTER = 230
COL_W = (PAGE_W - 2 * MARGIN_LR - GUTTER) // 2  # 5040 twips = 3.5 in
TEXT_W = PAGE_W - 2 * MARGIN_LR
FONT = "Times New Roman"


# ----------------------------------------------------------------------------- LaTeX helpers
def braced(s: str, i: int) -> tuple[str, int]:
    """s[i] == '{' -> (content, index after the matching '}')."""
    assert s[i] == "{", s[i:i + 30]
    depth, j = 0, i
    while True:
        c = s[j]
        if c == "\\":
            j += 2
            continue
        if c == "{":
            depth += 1
        elif c == "}":
            depth -= 1
            if depth == 0:
                return s[i + 1:j], j + 1
        j += 1


def read_macros(*files: Path) -> dict[str, str]:
    macros = {}
    for f in files:
        s = f.read_text(encoding="utf-8")
        for m in re.finditer(r"\\newcommand\{\\([A-Za-z]+)\}", s):
            val, _ = braced(s, m.end())
            macros[m.group(1)] = val.lstrip("%\n")
    return macros


def expand(s: str, macros: dict[str, str]) -> str:
    for _ in range(3):  # tables contain macros
        for name in sorted(macros, key=len, reverse=True):
            s = re.sub(r"\\" + name + r"(?![A-Za-z])(\{\})?", lambda _m, v=macros[name]: v, s)
    return s


def env_blocks(s: str, env: str):
    """Yield (start, end, inner, opt) for each \\begin{env}[opt] ... \\end{env} (non-nested)."""
    pat = re.compile(r"\\begin\{" + re.escape(env) + r"\}(\[[^\]]*\])?")
    pos = 0
    while True:
        m = pat.search(s, pos)
        if not m:
            return
        end_tag = "\\end{" + env + "}"
        e = s.index(end_tag, m.end())
        yield m.start(), e + len(end_tag), s[m.end():e], m.group(1)
        pos = e + len(end_tag)


def cite_text(nums: list[int]) -> str:
    nums = sorted(set(nums))
    groups, run = [], [nums[0]]
    for n in nums[1:]:
        if n == run[-1] + 1:
            run.append(n)
        else:
            groups.append(run)
            run = [n]
    groups.append(run)
    parts = []
    for g in groups:
        if len(g) >= 3:
            parts.append(f"[{g[0]}]--[{g[-1]}]")
        else:
            parts.extend(f"[{n}]" for n in g)
    return ", ".join(parts)


def extract_cmd(s: str, cmd: str) -> tuple[str | None, str]:
    """Remove the first \\cmd{...} from s; return (argument, s without it)."""
    i = s.find("\\" + cmd + "{")
    if i < 0:
        return None, s
    arg, j = braced(s, i + len(cmd) + 1)
    return arg, s[:i] + s[j:]


# ----------------------------------------------------------------------------- figures
def pdf_to_png(pdf: Path, png: Path, dpi=300):
    doc = pymupdf.open(pdf)
    pix = doc[0].get_pixmap(dpi=dpi)
    pix.save(png)
    return doc[0].rect.width / 72.0  # natural width in inches


def tikz_to_png(code: str, preamble: str, work: Path, name: str) -> tuple[Path, float]:
    tex = work / f"{name}.tex"
    tex.write_text("\\documentclass[border=3pt]{standalone}\n" + preamble + "\n\\begin{document}\n" + code +
                   "\n\\end{document}\n", encoding="utf-8")
    engine = os.environ.get("TECTONIC", "tectonic")
    subprocess.run([engine, tex.name], cwd=work, check=True, capture_output=True)
    png = work / f"{name}.png"
    return png, pdf_to_png(work / f"{name}.pdf", png, dpi=400)


# ----------------------------------------------------------------------------- preprocessing
def preprocess(build: Path, work: Path):
    src = (HERE / "faceguard.tex").read_text(encoding="utf-8")
    macros = read_macros(*sorted((HERE / "generated").glob("*.tex")))
    macros["sys"] = "FaceGuard"
    aux = (build / "faceguard.aux").read_text(encoding="utf-8")
    labels = {}
    for mm in re.finditer(r"\\newlabel\{([^}]*)\}\{", aux):
        outer, _ = braced(aux, mm.end() - 1)
        first, _ = braced(outer, 0)
        labels[mm.group(1)] = re.sub(r"\\mbox\s*\{([^}]*)\}", r"\1", first).strip()  # subsections: \mbox {IV-E}
    cites = {k: int(v) for k, v in re.findall(r"\\bibcite\{([^}]*)\}\{(\d+)\}", aux)}

    # TikZ preamble (colours and libraries) for standalone diagrams
    pre = src[:src.index("\\begin{document}")]
    colours = "\n".join(re.findall(r"\\definecolor.*", pre))
    tikz_pre = ("\\usepackage{amsmath}\\usepackage{newtxmath}\\usepackage[no-math]{fontspec}\n"
                "\\setmainfont{texgyretermes}[Extension=.otf, UprightFont=*-regular, BoldFont=*-bold, "
                "ItalicFont=*-italic, BoldItalicFont=*-bolditalic]\n\\usepackage{xcolor}\\usepackage{tikz}\n"
                "\\usetikzlibrary{arrows.meta,positioning,fit,calc,shapes.geometric,backgrounds}\n" + colours)

    body = src[src.index("\\begin{document}") + len("\\begin{document}"):src.index("\\end{document}")]
    body = expand(body, macros)
    body = re.sub(r"\\ref\{([^}]*)\}", lambda m: labels[m.group(1)], body)
    body = re.sub(r"\\cite\{([^}]*)\}", lambda m: cite_text([cites[k.strip()] for k in m.group(1).split(",")]), body)

    # front matter
    head = src[:src.index("\\begin{document}")] + body[:body.index("\\maketitle")]
    head = expand(head, macros)
    title, _ = braced(head, head.index("\\title{") + 6)
    author, _ = braced(head, head.index("\\author{") + 7)
    thanks = []
    while True:
        t, author = extract_cmd(author, "thanks")
        if t is None:
            break
        thanks.append(t.strip().rstrip("%"))
    author = re.sub(r"\\IEEEmembership\{([^}]*)\}", r"\\emph{\1}", author).replace("%", "").strip()
    markboth, _ = braced(head, head.index("\\markboth{") + 9)
    header = re.sub(r"~|\\,", " ", markboth).replace("%", "").strip()
    abstract = next(env_blocks(body, "abstract"))[2].strip()
    keywords = next(env_blocks(body, "IEEEkeywords"))[2].strip()
    rest = body[body.index("\\end{IEEEkeywords}") + len("\\end{IEEEkeywords}"):]
    # The first-page note is attached to the abstract (two-column section): a footnote in the single-column
    # title section would push the body to page 2 in Word.
    front = (f"@@TITLE@@ {title}\n\n@@AUTHORS@@ {author}\n\n"
             f"@@ABSTRACT@@ \\footnote{{" + "\n\n".join(thanks) + f"}}{abstract}\n\n@@KEYWORDS@@ {keywords}\n\n")

    rest = re.sub(r"\\IEEEPARstart\{(\w)\}\{(\w+)\}", lambda m: "@@DROPCAP@@" + m.group(1) + m.group(2), rest)
    rest = rest.replace("\\mbox{0.08}", "0.08")
    rest = re.sub(r"\\label\{[^}]*\}", "", rest)
    rest = re.sub(r"\\IEEEtriggeratref\{[^}]*\}", "", rest)
    rest = re.sub(r"\\cmidrule(\([^)]*\))?\{[^}]*\}", "", rest)

    # algorithms -> caption + numbered pseudocode lines
    out, pos, n_alg = [], 0, 0
    for a, b, inner, _ in env_blocks(rest, "algorithm"):
        n_alg += 1
        out.append(rest[pos:a])
        out.append(convert_algorithm(inner, n_alg))
        pos = b
    out.append(rest[pos:])
    rest = "".join(out)

    # figures -> placeholders (images handled after pandoc)
    figs = {}
    out, pos = [], 0
    matches = sorted([(a, b, inner, "*") for a, b, inner, _ in env_blocks(rest, "figure*")] +
                     [(a, b, inner, "") for a, b, inner, _ in env_blocks(rest, "figure")])
    fig_no = 0
    for a, b, inner, star in matches:
        fig_no += 1
        caption, _ = braced(inner, inner.index("\\caption{") + 8)
        images = []
        if "\\begin{tikzpicture}" in inner:
            code = inner[inner.index("\\begin{tikzpicture}"):inner.index("\\end{tikzpicture}") + len("\\end{tikzpicture}")]
            png, w = tikz_to_png(code, tikz_pre, work, f"tikz{fig_no}")
            images.append((png, w))
        for g in re.findall(r"\\includegraphics(?:\[[^\]]*\])?\{([^}]*)\}", inner):
            f = HERE / g
            if f.suffix.lower() == ".pdf":
                png = work / (f.stem + ".png")
                images.append((png, pdf_to_png(f, png)))
            else:
                images.append((f, None))
        figs[fig_no] = {"images": images, "wide": bool(star)}
        out.append(rest[pos:a])
        out.append(f"\n\n@@FIG:{fig_no}@@\n\n@@FIGCAP:{fig_no}@@ {caption}\n\n")
        pos = b
    out.append(rest[pos:])
    rest = "".join(out)

    # tables -> caption marker + tabular (+ note); table* (full width) recorded in wide_tables
    tables, wide_tables = [], set()
    out, pos = [], 0
    tmatches = sorted([(a, b, inner, "*") for a, b, inner, _ in env_blocks(rest, "table*")] +
                      [(a, b, inner, "") for a, b, inner, _ in env_blocks(rest, "table")])
    for a, b, inner, star in tmatches:
        if star:
            wide_tables.add(len(tables) + 1)
        caption, _ = braced(inner, inner.index("\\caption{") + 8)
        t0 = inner.index("\\begin{tabular}")
        t1 = inner.index("\\end{tabular}") + len("\\end{tabular}")
        spec, after_spec = braced(inner, t0 + len("\\begin{tabular}"))
        spec = spec.replace("@{}", "")
        spec = re.sub(r">\{[^{}]*\}", "", spec)  # pandoc mis-parses array-package column prefixes (\raggedright etc.)
        rows = inner[after_spec:inner.index("\\end{tabular}")]
        note = inner[t1:].replace("\\\\[2pt]", "").replace("\\raggedright", "").replace("\\scriptsize", "").strip()
        rows = rows.replace("\\toprule", "").replace("\\bottomrule", "")
        rows = re.sub(r"\\multicolumn\{(\d+)\}\{@\{\}l\}", r"\\multicolumn{\1}{l}", rows)  # pandoc rejects @{} here
        rows = rows.replace("\\quad ", "\u2003").replace("\\quad", "\u2003")
        # first \midrule ends the header; later ones become a marker on the next row
        first = rows.index("\\midrule")
        head_rows, body_rows = rows[:first], rows[first + len("\\midrule"):]
        # mark the last header row (pandoc keeps only the first row as a Word header row)
        hr = head_rows.split("\\\\")
        last = max(i for i, r in enumerate(hr) if r.strip())
        hr[last] = hr[last].replace(hr[last].lstrip(), "@@HE@@" + hr[last].lstrip(), 1)
        head_rows = "\\\\".join(hr)
        body_rows = re.sub(r"\\midrule\s*", "@@MR@@", body_rows)
        # pandoc only parses \multicolumn at the start of a cell, so the marker goes inside it
        body_rows = re.sub(r"@@MR@@(\\multicolumn\{\d+\}\{[^}]*\}\{)", r"\1@@MR@@", body_rows)
        tables.append(caption)
        out.append(rest[pos:a])
        out.append(f"\n\n@@TABCAP:{len(tables)}@@ {caption}\n\n\\begin{{tabular}}{{{spec}}}\n\\toprule\n{head_rows}\\midrule\n"
                   f"{body_rows}\\bottomrule\n\\end{{tabular}}\n\n" + (f"@@TABNOTE@@ {note}\n\n" if note else ""))
        pos = b
    out.append(rest[pos:])
    rest = "".join(out)

    # numbered equations -> marker paragraphs
    out, pos, n = [], 0, 0
    for a, b, inner, _ in env_blocks(rest, "equation"):
        n += 1
        follow = rest[b:b + 3]
        cont = not follow.startswith("\n\n")
        out.append(rest[pos:a])
        out.append(f"\n\n@@EQ:{n}@@\n\n\\[{inner.strip()}\\]\n\n" + ("@@NOINDENT@@" if cont else ""))
        pos = b
        if cont:
            while pos < len(rest) and rest[pos] in " \n":
                pos += 1
    out.append(rest[pos:])
    rest = "".join(out)

    # theorem-like environments
    for a, b, inner, _ in reversed(list(env_blocks(rest, "proposition"))):
        rest = rest[:a] + f"\\textbf{{Proposition 1.}} \\emph{{{inner.strip()}}}" + rest[b:]
    for a, b, inner, _ in reversed(list(env_blocks(rest, "proof"))):
        rest = rest[:a] + f"\n\n\\emph{{Proof.}} {inner.strip()}@@QED@@\n\n" + rest[b:]

    # references from the .bbl (already in citation order, IEEEtran formatting)
    bbl = (build / "faceguard.bbl").read_text(encoding="utf-8")
    items = re.split(r"\\bibitem\{[^}]*\}", bbl[:bbl.index("\\end{thebibliography}")])[1:]
    refs = []
    for k, it in enumerate(items, 1):
        it = re.sub(r"\\BIBentry(ALT|STD)interwordspacing\s*", "", it).replace("\\newblock", "")
        it = re.sub(r"\\hskip[^\\]*\\relax", " ", it)  # IEEEtran's publisher spacing
        refs.append(f"@@REF:{k}@@ " + " ".join(it.split()))
    bib = "\\section*{References}\n\n" + "\n\n".join(refs) + "\n\n"
    # LaTeX puts a thin space after commas in math; Word's math engine does not
    def comma_space(m):
        parts = re.split(r"(\\text\{[^}]*\})", m.group(0))  # leave text-mode commas alone
        return "".join(s if s.startswith("\\text{") else re.sub(r",(?!\\,)", r",\\,", s) for s in parts)
    rest = re.sub(r"\$[^$]+\$", comma_space, rest)
    rest = re.sub(r"\\\[.*?\\\]", comma_space, rest, flags=re.S)
    front = re.sub(r"\$[^$]+\$", comma_space, front)

    rest = re.sub(r"\\bibliographystyle\{[^}]*\}", "", rest)
    rest = rest.replace("\\bibliography{references}", bib)
    rest = rest.replace("\\balance", "")

    doc_tex = "\\documentclass{article}\n\\usepackage{amsmath}\n\\begin{document}\n" + front + rest + "\n\\end{document}\n"
    (work / "pre.tex").write_text(doc_tex, encoding="utf-8")
    return figs, header, wide_tables


def convert_algorithm(inner: str, n: int) -> str:
    """algpseudocode -> marker paragraphs: '@@ALGCAP:n@@ caption' and '@@ALGLINE:level:number@@ text'."""
    cap, _ = braced(inner, inner.index("\\caption{") + 8)
    body = inner[inner.index("\\begin{algorithmic}"):inner.index("\\end{algorithmic}")]
    body = re.sub(r"\\begin\{algorithmic\}(\[[^\]]*\])?", "", body)
    kw = lambda w: f"\\textbf{{{w}}}"  # noqa: E731
    lines, level, num = [], 0, 0

    def emit(text, numbered=True):
        nonlocal num
        if numbered:
            num += 1
        lines.append(f"@@ALGLINE:{level}:{num if numbered else 0}@@ {text.strip()}")

    def arg(s, cmd):
        i = s.index("\\" + cmd) + len(cmd) + 1
        a1, j = braced(s, i)
        return a1, s[j:].strip()

    for raw in body.splitlines():
        s = raw.strip().replace("\\Return", kw("return"))
        if not s or s.startswith("\\footnotesize"):
            continue
        if s.startswith("\\Require"):
            emit(kw("Input:") + " " + s[len("\\Require"):], numbered=False)
        elif s.startswith("\\Ensure"):
            emit(kw("Output:") + " " + s[len("\\Ensure"):], numbered=False)
        elif s.startswith("\\IfThen"):
            c, rest_ = arg(s, "IfThen")
            a2, _ = braced(rest_, 0)
            emit(f"{kw('if')} {c} {kw('then')} {a2}")
        elif s.startswith("\\ElsIf"):
            c, tail = arg(s, "ElsIf")
            level -= 1
            emit(f"{kw('else if')} {c} {kw('then')} {tail}")
            level += 1
        elif s.startswith("\\Else"):
            level -= 1
            emit(kw("else"))
            level += 1
        elif s.startswith("\\If"):
            c, tail = arg(s, "If")
            emit(f"{kw('if')} {c} {kw('then')} {tail}")
            level += 1
        elif s.startswith("\\EndIf"):
            level -= 1
            emit(kw("end if"))
        elif s.startswith("\\For"):
            c, tail = arg(s, "For")
            emit(f"{kw('for')} {c} {kw('do')} {tail}")
            level += 1
        elif s.startswith("\\EndFor"):
            level -= 1
            emit(kw("end for"))
        elif s.startswith("\\Repeat"):
            emit(kw("repeat"))
            level += 1
        elif s.startswith("\\Until"):
            c, _ = arg(s, "Until")
            level -= 1
            emit(f"{kw('until')} {c}")
        elif s.startswith("\\State"):
            emit(s[len("\\State"):])
        else:  # continuation of the previous statement
            lines[-1] += " " + s
    return f"\n\n@@ALGCAP:{n}@@ \\textbf{{Algorithm {n}}} {cap}\n\n" + "\n\n".join(lines) + "\n\n"


# ----------------------------------------------------------------------------- docx helpers
def set_fonts(rpr, size=None, bold=None, italic=None, smallcaps=None):
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(att), None)
    for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(att), FONT)


def style_font(style, size, bold=False, italic=False, smallcaps=False, color_black=True):
    f = style.font
    f.name = FONT
    set_fonts(style.element.get_or_add_rPr())
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.small_caps = smallcaps
    if color_black:
        f.color.rgb = RGBColor(0, 0, 0)


def para_fmt(style, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first=0.0, left=0.0, before=0, after=0, keep=False, hanging=None):
    pf = style.paragraph_format
    pf.alignment = align
    pf.first_line_indent = Inches(first) if hanging is None else -Inches(hanging)
    pf.left_indent = Inches(left)
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.keep_with_next = keep
    pf.widow_control = True


def get_style(doc, name, base=None):
    try:
        return doc.styles[name]
    except KeyError:
        st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        if base:
            st.base_style = doc.styles[base]
        return st


def para_text(p) -> str:
    return "".join(t.text or "" for t in p._p.iter(qn("w:t")))


def strip_marker(p, marker: str) -> None:
    """Remove the first occurrence of marker (which may span runs) from paragraph p."""
    ts = list(p._p.iter(qn("w:t")))
    full = "".join(t.text or "" for t in ts)
    i = full.find(marker)
    if i < 0:
        return
    j = i + len(marker)
    # also swallow one following space
    if j < len(full) and full[j] == " ":
        j += 1
    pos = 0
    for t in ts:
        txt = t.text or ""
        a, b = pos, pos + len(txt)
        lo, hi = max(a, i), min(b, j)
        if lo < hi:
            t.text = txt[:lo - a] + txt[hi - a:]
            t.set(qn("xml:space"), "preserve")
        pos = b


def insert_run_first(p, text, bold=False, italic=False):
    r = OxmlElement("w:r")
    if bold or italic:
        rpr = OxmlElement("w:rPr")
        if bold:
            rpr.append(OxmlElement("w:b"))
        if italic:
            rpr.append(OxmlElement("w:i"))
        r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    ppr = p._p.find(qn("w:pPr"))
    if ppr is not None:
        ppr.addnext(r)
    else:
        p._p.insert(0, r)
    return r


def sect_pr(cols: int, continuous: bool, header_ref=None):
    s = OxmlElement("w:sectPr")
    if header_ref is not None:
        s.append(copy.deepcopy(header_ref))
    if continuous:
        t = OxmlElement("w:type")
        t.set(qn("w:val"), "continuous")
        s.append(t)
    sz = OxmlElement("w:pgSz")
    sz.set(qn("w:w"), str(PAGE_W))
    sz.set(qn("w:h"), str(PAGE_H))
    s.append(sz)
    mar = OxmlElement("w:pgMar")
    for k, v in (("top", MARGIN_T), ("right", MARGIN_LR), ("bottom", MARGIN_B), ("left", MARGIN_LR),
                 ("header", 580), ("footer", 580), ("gutter", 0)):
        mar.set(qn("w:" + k), str(v))
    s.append(mar)
    c = OxmlElement("w:cols")
    c.set(qn("w:num"), str(cols))
    c.set(qn("w:space"), str(GUTTER))
    s.append(c)
    return s


def end_section_at(p, cols, continuous=True):
    ppr = p._p.get_or_add_pPr()
    old = ppr.find(qn("w:sectPr"))
    if old is not None:
        ppr.remove(old)
    ppr.append(sect_pr(cols, continuous))


def tc_border(cell, side, sz=4):
    tcpr = cell._tc.get_or_add_tcPr()
    b = tcpr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcpr.append(b)
    e = OxmlElement("w:" + side)
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), "0")
    e.set(qn("w:color"), "000000")
    b.append(e)


def table_borders(tbl, top=8, bottom=8):
    tblpr = tbl._tbl.tblPr
    for tag in ("w:tblBorders", "w:tblStyle", "w:tblW", "w:tblLayout", "w:jc", "w:tblCellMar", "w:tblLook"):
        for e in tblpr.findall(qn(tag)):
            tblpr.remove(e)
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), "0")
    w.set(qn("w:type"), "auto")
    tblpr.append(w)
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "center")
    tblpr.append(jc)
    borders = OxmlElement("w:tblBorders")
    for side, sz in (("top", top), ("left", 0), ("bottom", bottom), ("right", 0), ("insideH", 0), ("insideV", 0)):
        e = OxmlElement("w:" + side)
        e.set(qn("w:val"), "single" if sz else "nil")
        if sz:
            e.set(qn("w:sz"), str(sz))
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), "000000")
        borders.append(e)
    tblpr.append(borders)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "autofit")
    tblpr.append(lay)
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", 0), ("left", 60), ("bottom", 0), ("right", 60)):
        e = OxmlElement("w:" + side)
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblpr.append(mar)
    for tcw in tbl._tbl.iter(qn("w:tcW")):
        tcw.set(qn("w:type"), "auto")
        tcw.set(qn("w:w"), "0")


# ----------------------------------------------------------------------------- post-processing
def postprocess(raw: Path, out: Path, figs: dict, header_text: str, wide_tables: set):
    d = docx.Document(str(raw))
    S = d.styles

    # document defaults: Times New Roman 10 pt
    rpr_default = d.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    set_fonts(rpr_default)
    sz = rpr_default.find(qn("w:sz"))
    if sz is None:
        sz = OxmlElement("w:sz")
        rpr_default.append(sz)
    sz.set(qn("w:val"), "20")

    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in [s.name for s in S]:
            style_font(S[name], 10)
            para_fmt(S[name], first=0.14 if name != "Compact" else 0.0)
    for name in ("Title", "Author", "Abstract", "Heading 1", "Heading 2", "Heading 3", "Footnote Text", "Hyperlink",
                 "Verbatim Char"):
        try:
            S[name]
        except KeyError:
            continue
    h1, h2 = S["Heading 1"], S["Heading 2"]
    style_font(h1, 10, smallcaps=True)
    para_fmt(h1, align=WD_ALIGN_PARAGRAPH.CENTER, before=10, after=4, keep=True)
    style_font(h2, 10, italic=True)
    para_fmt(h2, align=WD_ALIGN_PARAGRAPH.LEFT, before=7, after=3, keep=True)
    for h in (h1, h2):
        h.paragraph_format.first_line_indent = Inches(0)
    fn = get_style(d, "Footnote Text")
    style_font(fn, 8)
    para_fmt(fn, first=0.12)
    try:
        hl = d.styles["Hyperlink"]
        hl.font.color.rgb = RGBColor(0, 0, 0)
        hl.font.underline = False
    except KeyError:
        pass
    try:
        vc = d.styles["Verbatim Char"]
        vc.font.name = "Courier New"
        vc.font.size = Pt(9)
    except KeyError:
        pass

    st_title = get_style(d, "IEEE Title")
    style_font(st_title, 24)
    para_fmt(st_title, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    st_auth = get_style(d, "IEEE Authors")
    style_font(st_auth, 11)
    para_fmt(st_auth, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
    st_abs = get_style(d, "IEEE Abstract")
    style_font(st_abs, 9, bold=True)
    para_fmt(st_abs, first=0.14, after=6)
    st_figcap = get_style(d, "IEEE Figure Caption")
    style_font(st_figcap, 8)
    para_fmt(st_figcap, before=3, after=8)
    st_tabcap = get_style(d, "IEEE Table Caption")
    style_font(st_tabcap, 8)
    para_fmt(st_tabcap, align=WD_ALIGN_PARAGRAPH.CENTER, before=8, after=3, keep=True)
    st_tab = get_style(d, "IEEE Table Text")
    style_font(st_tab, 8)
    para_fmt(st_tab, align=WD_ALIGN_PARAGRAPH.LEFT)
    st_note = get_style(d, "IEEE Table Note")
    style_font(st_note, 7)
    para_fmt(st_note, align=WD_ALIGN_PARAGRAPH.LEFT, before=2, after=8)
    st_fig = get_style(d, "IEEE Figure")
    style_font(st_fig, 10)
    para_fmt(st_fig, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, keep=True)
    st_ref = get_style(d, "IEEE Reference")
    style_font(st_ref, 8)
    para_fmt(st_ref, left=0.3, hanging=0.3)
    st_ref.paragraph_format.tab_stops.add_tab_stop(Inches(0.3))
    st_eq = get_style(d, "IEEE Equation")
    style_font(st_eq, 10)
    para_fmt(st_eq, align=WD_ALIGN_PARAGRAPH.CENTER)

    # heading numbering: I. / A. (multilevel list linked to the heading styles)
    numbering = d.part.numbering_part.element
    abs_ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    aid, nid = max(abs_ids + [0]) + 1, max(num_ids + [0]) + 1
    an = OxmlElement("w:abstractNum")
    an.set(qn("w:abstractNumId"), str(aid))
    mlt = OxmlElement("w:multiLevelType")
    mlt.set(qn("w:val"), "multilevel")
    an.append(mlt)
    for lvl, fmt, text in ((0, "upperRoman", "%1."), (1, "upperLetter", "%2.")):
        L = OxmlElement("w:lvl")
        L.set(qn("w:ilvl"), str(lvl))
        for tag, val in (("w:start", "1"), ("w:numFmt", fmt), ("w:pStyle", f"Heading{lvl + 1}"),
                         ("w:suff", "space"), ("w:lvlText", text), ("w:lvlJc", "left")):
            e = OxmlElement(tag)
            e.set(qn("w:val"), val)
            L.append(e)
        ppr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "0")
        ind.set(qn("w:firstLine"), "0")
        ppr.append(ind)
        L.append(ppr)
        an.append(L)
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(an)
    else:
        numbering.append(an)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(nid))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(aid))
    num.append(ref)
    numbering.append(num)
    for style, lvl in ((h1, 0), (h2, 1)):
        ppr = style.element.get_or_add_pPr()
        np_ = OxmlElement("w:numPr")
        il = OxmlElement("w:ilvl")
        il.set(qn("w:val"), str(lvl))
        ni = OxmlElement("w:numId")
        ni.set(qn("w:val"), str(nid))
        np_.append(il)
        np_.append(ni)
        ppr.insert(0, np_)
    # enumerate lists: "1)" like IEEEtran
    for lvl in numbering.iter(qn("w:lvl")):
        fmt = lvl.find(qn("w:numFmt"))
        txt = lvl.find(qn("w:lvlText"))
        if fmt is not None and txt is not None and fmt.get(qn("w:val")) == "decimal" and txt.get(qn("w:val")) == "%1.":
            txt.set(qn("w:val"), "%1)")

    body = d.element.body
    paras = list(d.paragraphs)
    continuation = []  # paragraphs that continue after an equation
    tabcap_p = {}  # table number -> caption paragraph
    st_alg = get_style(d, "IEEE Algorithm")
    style_font(st_alg, 8.5)
    para_fmt(st_alg, align=WD_ALIGN_PARAGRAPH.LEFT)
    alg_last = {}  # algorithm number -> last line paragraph
    cur_alg = 0
    in_refs = False
    title_p = authors_p = None
    for p in paras:
        t = para_text(p)
        sname = p.style.name if p.style is not None else ""
        if sname.startswith("Heading"):
            txt = t.strip()
            if txt in ("Acknowledgment", "References"):
                ppr = p._p.get_or_add_pPr()
                np_ = OxmlElement("w:numPr")
                ni = OxmlElement("w:numId")
                ni.set(qn("w:val"), "0")
                np_.append(ni)
                ppr.insert(0, np_)
            in_refs = txt == "References"
            continue
        if t.startswith("@@TITLE@@"):
            strip_marker(p, "@@TITLE@@")
            p.style = st_title
            title_p = p
        elif t.startswith("@@AUTHORS@@"):
            strip_marker(p, "@@AUTHORS@@")
            p.style = st_auth
            authors_p = p
        elif t.startswith("@@ABSTRACT@@"):
            strip_marker(p, "@@ABSTRACT@@")
            p.style = st_abs
            insert_run_first(p, "Abstract\u2014", bold=True, italic=True)
        elif t.startswith("@@KEYWORDS@@"):
            strip_marker(p, "@@KEYWORDS@@")
            p.style = st_abs
            insert_run_first(p, "Index Terms\u2014", bold=True, italic=True)
            p.paragraph_format.space_after = Pt(10)
        elif t.startswith("@@TABCAP:"):
            n = int(re.match(r"@@TABCAP:(\d+)@@", t).group(1))
            strip_marker(p, f"@@TABCAP:{n}@@")
            p.style = st_tabcap
            for r in p.runs:
                r.font.small_caps = True
            r0 = insert_run_first(p, f"TABLE {ROMAN[n]}")
            br = OxmlElement("w:br")
            r0.append(br)
            tabcap_p[n] = p
        elif t.startswith("@@ALGCAP:"):
            n = int(re.match(r"@@ALGCAP:(\d+)@@", t).group(1))
            strip_marker(p, f"@@ALGCAP:{n}@@")
            p.style = st_alg
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.keep_with_next = True
            border_para(p, ("top", 12), ("bottom", 6))
            cur_alg = n
        elif t.startswith("@@ALGLINE:"):
            m_ = re.match(r"@@ALGLINE:(-?\d+):(\d+)@@", t)
            lvl, num = int(m_.group(1)), int(m_.group(2))
            strip_marker(p, m_.group(0))
            p.style = st_alg
            hang = 0.24
            p.paragraph_format.left_indent = Inches(hang + 0.17 * max(0, lvl))
            p.paragraph_format.first_line_indent = -Inches(hang + 0.17 * max(0, lvl))
            p.paragraph_format.tab_stops.add_tab_stop(Inches(hang + 0.17 * max(0, lvl)))
            p.paragraph_format.keep_with_next = True
            r = insert_run_first(p, f"{num}:" if num else "")
            r.append(OxmlElement("w:tab"))
            alg_last[cur_alg] = p
        elif t.startswith("@@TABNOTE@@"):
            strip_marker(p, "@@TABNOTE@@")
            p.style = st_note
        elif t.startswith("@@FIGCAP:"):
            n = int(re.match(r"@@FIGCAP:(\d+)@@", t).group(1))
            strip_marker(p, f"@@FIGCAP:{n}@@")
            p.style = st_figcap
            insert_run_first(p, f"Fig. {n}. ")
            if figs[n]["wide"]:
                p.paragraph_format.space_after = Pt(10)
        elif t.startswith("@@FIG:"):
            n = int(re.match(r"@@FIG:(\d+)@@", t).group(1))
            for r in list(p._p.findall(qn("w:r"))):
                p._p.remove(r)
            p.style = st_fig
            imgs = figs[n]["images"]
            maxw = TEXT_W / 1440 if figs[n]["wide"] else COL_W / 1440
            for k, (path, natural) in enumerate(imgs):
                w = maxw if natural is None else min(maxw, natural)
                run = p.add_run()
                run.add_picture(str(path), width=Inches(w))
                if k < len(imgs) - 1:
                    run.add_break()
        elif t.startswith("@@NOINDENT@@"):
            strip_marker(p, "@@NOINDENT@@")
            p.paragraph_format.first_line_indent = Inches(0)
            continuation.append(p._p)  # keep references: lxml proxies are then stable
        elif t.startswith("@@DROPCAP@@"):
            strip_marker(p, "@@DROPCAP@@")
            drop_cap(p)
        elif t.startswith("@@REF:"):
            n = int(re.match(r"@@REF:(\d+)@@", t).group(1))
            strip_marker(p, f"@@REF:{n}@@")
            p.style = st_ref
            r = insert_run_first(p, f"[{n}]")
            tab = OxmlElement("w:tab")
            r.append(tab)
        if "@@QED@@" in t:
            strip_marker(p, "@@QED@@")
            p.paragraph_format.tab_stops.add_tab_stop(Twips(COL_W), WD_TAB_ALIGNMENT.RIGHT)
            run = p.add_run()
            run._r.append(OxmlElement("w:tab"))
            run.add_text("\u25a1")

    # equations: marker paragraph + math paragraph -> borderless 3-cell table with the number
    paras = list(d.paragraphs)
    for i, p in enumerate(paras):
        t = para_text(p)
        m = re.match(r"@@EQ:(\d+)@@", t)
        if not m:
            continue
        n = int(m.group(1))
        omath = math_p = None
        for cand in (p, paras[i + 1]):
            omath = next(cand._p.iter(qn("m:oMathPara")), None)
            if omath is None:
                omath = next(cand._p.iter(qn("m:oMath")), None)
            if omath is not None:
                math_p = cand
                break
        if omath is None:
            raise RuntimeError(f"equation {n}: no math found; next paragraph text: {para_text(paras[i + 1])[:120]!r}")
        tbl = d.add_table(rows=1, cols=3)
        widths = (560, COL_W - 1120, 560)
        for c, w in zip(tbl.rows[0].cells, widths):
            c.width = Twips(w)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            c.paragraphs[0].style = st_eq
        tbl.rows[0].cells[1].paragraphs[0]._p.append(omath)
        num_p = tbl.rows[0].cells[2].paragraphs[0]
        num_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        num_p.add_run(f"({n})")
        tblpr = tbl._tbl.tblPr
        for e in tblpr.findall(qn("w:tblStyle")) + tblpr.findall(qn("w:tblLook")):
            tblpr.remove(e)
        cap = OxmlElement("w:tblCaption")
        cap.set(qn("w:val"), f"Equation {n}")
        tblpr.append(cap)
        for c in tbl.rows[0].cells:
            tcpr = c._tc.get_or_add_tcPr()
            nb = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                e = OxmlElement("w:" + side)
                e.set(qn("w:val"), "nil")
                nb.append(e)
            tcpr.append(nb)
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement("w:" + side)
            e.set(qn("w:val"), "nil")
            borders.append(e)
        tblpr.append(borders)
        mar = OxmlElement("w:tblCellMar")
        for side in ("left", "right"):
            e = OxmlElement("w:" + side)
            e.set(qn("w:w"), "0")
            e.set(qn("w:type"), "dxa")
            mar.append(e)
        tblpr.append(mar)
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        tblpr.append(lay)
        p._p.addprevious(tbl._tbl)
        for q in {id(p._p): p._p, id(math_p._p): math_p._p}.values():
            q.getparent().remove(q)

    # data tables: IEEE rules, 8-pt text, centred, header / midrule markers
    for tbl in d.tables:
        cap = tbl._tbl.tblPr.find(qn("w:tblCaption"))
        if cap is not None and cap.get(qn("w:val"), "").startswith("Equation"):
            continue
        table_borders(tbl)
        for row in tbl.rows:
            for cell in row.cells:
                for cp in cell.paragraphs:
                    jc = cp.paragraph_format.alignment
                    cp.style = st_tab
                    if jc is not None:
                        cp.alignment = jc
                    for r in cp.runs:
                        r.font.size = None
            for marker, side in (("@@HE@@", "bottom"), ("@@MR@@", "top")):
                if marker in "".join(para_text(cp) for c in row.cells for cp in c.paragraphs):
                    for cell in row.cells:
                        for cp in cell.paragraphs:
                            if marker in para_text(cp):
                                strip_marker(cp, marker)
                        tc_border(cell, side, 4)
        # never split a table across columns or pages (IEEE style)
        for ri, row in enumerate(tbl.rows):
            trpr = row._tr.get_or_add_trPr()
            if trpr.find(qn("w:cantSplit")) is None:
                trpr.insert(0, OxmlElement("w:cantSplit"))
            if ri < len(tbl.rows) - 1:
                for cell in row.cells:
                    for cp in cell.paragraphs:
                        cp.paragraph_format.keep_with_next = True
        # breathing space between a table and the text below it
        nxt = tbl._tbl.getnext()
        if nxt is not None and nxt.tag == qn("w:p"):
            np_ = docx.text.paragraph.Paragraph(nxt, tbl._parent)
            if np_.style is not None and np_.style.name in ("Body Text", "First Paragraph", "IEEE Figure Caption"):
                np_.paragraph_format.space_before = Pt(8)
            elif np_.style is not None and np_.style.name == "IEEE Table Caption":
                pass
    # compact list indentation (bullets and "1)" lists), leaving the heading numbering alone
    for an_el in numbering.findall(qn("w:abstractNum")):
        if an_el is an:
            continue
        for lvl in an_el.findall(qn("w:lvl")):
            k = int(lvl.get(qn("w:ilvl")))
            ind = lvl.find(qn("w:pPr") + "/" + qn("w:ind"))
            if ind is not None:
                ind.set(qn("w:left"), str(360 + 360 * k))
                ind.set(qn("w:hanging"), "260")

    # algorithm blocks end with a rule, like the algorithm package
    for p in alg_last.values():
        border_para(p, ("bottom", 6))
        p.paragraph_format.keep_with_next = False
        p.paragraph_format.space_after = Pt(8)

    # full-width tables float at the top of the page like the wide figures
    for n in sorted(wide_tables):
        cap = tabcap_p[n]
        tbl = cap._p.getnext()
        parts = [cap._p, tbl]
        nxt = tbl.getnext()
        if nxt is not None and nxt.tag == qn("w:p") and \
                docx.text.paragraph.Paragraph(nxt, cap._parent).style.name == "IEEE Table Note":
            parts.append(nxt)
        rows = len(tbl.findall(qn("w:tr")))
        height = int((rows * 11.0 + 24 + (12 if len(parts) > 2 else 0) + 10) * EMU_PER_PT)
        anchor = cap._p.getprevious()
        while anchor is not None and anchor.tag != qn("w:p"):
            anchor = anchor.getprevious()
        float_box(parts, anchor, 100 + n, height, f"Table {n}")

    # section layout: title block single-column, body two-column.
    # Wide figures float at the top of the page (like figure* in LaTeX): figure + caption go into a
    # full-width text box with top-and-bottom wrapping, anchored to the next body paragraph.
    end_section_at(authors_p, 1, continuous=False)
    body_styles = ("Body Text", "First Paragraph")
    paras = list(d.paragraphs)
    for i, p in enumerate(paras):
        if p.style is None or p.style.name != "IEEE Figure":
            continue
        cap = paras[i + 1]
        n = int(re.match(r"Fig\. (\d+)\.", para_text(cap)).group(1))
        if not figs[n]["wide"]:
            continue
        # Anchor to the paragraph just before the figure's place in the text (usually its section heading):
        # an anchor that sits later, near a page bottom, makes Word bounce the float between pages and leave gaps.
        anchor = p._p.getprevious()
        while anchor is not None and anchor.tag != qn("w:p"):
            anchor = anchor.getprevious()
        float_top(p, cap, anchor, n)
    # The last section (after the title block) is two-column and ends with the last reference. No balancing
    # paragraph is appended: when the last page is full, Word would push that empty paragraph onto a blank page.
    final = d.sections[-1]._sectPr
    for e in list(final):
        final.remove(e)
    for e in sect_pr(2, True):
        final.append(e)

    # running header on every page
    sec0 = d.sections[0]
    sec0.header.is_linked_to_previous = False
    hp = sec0.header.paragraphs[0]
    for r in list(hp.runs):
        r._r.getparent().remove(r._r)
    hp.paragraph_format.tab_stops.add_tab_stop(Twips(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
    run = hp.add_run(header_text.upper())
    run.font.size = Pt(7)
    run.font.name = FONT
    run = hp.add_run()
    run._r.append(OxmlElement("w:tab"))
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    s_ = OxmlElement("w:sz")
    s_.set(qn("w:val"), "16")
    rpr.append(s_)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    hp._p.append(fld)

    # first-page note (the first footnote): no visible mark; later footnotes keep their numbers
    first_ref = next(body.iter(qn("w:footnoteReference")), None)
    first_id = None
    if first_ref is not None:
        first_ref.set(qn("w:customMarkFollows"), "1")
        first_id = first_ref.get(qn("w:id"))
    fn_part = next((rel.target_part for rel in d.part.rels.values() if rel.reltype.endswith("/footnotes")), None)
    if fn_part is not None and first_id is not None:
        from lxml import etree
        root = etree.fromstring(fn_part.blob)
        target = next(f for f in root.findall(qn("w:footnote")) if f.get(qn("w:id")) == first_id)
        for ref_run in list(target.iter(qn("w:footnoteRef"))):
            run_el = ref_run.getparent()
            nxt = run_el.getnext()
            run_el.getparent().remove(run_el)
            if nxt is not None and nxt.tag == qn("w:r") and "".join(nxt.itertext()).strip() == "":
                nxt.getparent().remove(nxt)
        fn_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    for root in (d.element, d.styles.element, d.part.numbering_part.element, sec0.header._element):
        fix_order(root)
    d.core_properties.title = para_text(title_p)
    d.core_properties.author = "FaceGuard authors"
    d.save(str(out))


EMU_PER_PT = 12700
WPS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"


def border_para(p, *sides):
    """Add paragraph borders, e.g. border_para(p, ("top", 12), ("bottom", 6)) (sizes in eighths of a point)."""
    ppr = p._p.get_or_add_pPr()
    bdr = ppr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        ppr.append(bdr)
    for side, sz in sides:
        e = OxmlElement("w:" + side)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "1")
        e.set(qn("w:color"), "000000")
        bdr.append(e)


def float_top(fig_p, cap_p, anchor_el, n):
    """Move a figure paragraph and its caption into a full-width text box floating at the top of the page."""
    img_h = sum(int(e.get("cy")) for e in fig_p._p.iter("{%s}extent" % WP_NS))
    chars = len(para_text(cap_p))
    lines = -(-chars // 150)  # 8-pt Times across 7.16 in
    height = img_h + int((lines * 9.4 + 3 + 6 + 8) * EMU_PER_PT)
    float_box([fig_p._p, cap_p._p], anchor_el, n, height, f"Figure {n}")


def float_box(elements, anchor_el, n, height, name):
    """Move body elements (paragraphs/tables) into a full-width text box floating at the top of the page."""
    from lxml import etree
    width = TEXT_W * 635  # twips -> EMU
    xml = f"""
<w:r xmlns:w="{W_NS[1:-1]}" xmlns:wp="{WP_NS}" xmlns:a="{A_NS}" xmlns:wps="{WPS}">
  <w:drawing>
    <wp:anchor distT="0" distB="{int(8 * EMU_PER_PT)}" distL="0" distR="0" simplePos="0" relativeHeight="{251659264 + n}"
               behindDoc="0" locked="0" layoutInCell="1" allowOverlap="0">
      <wp:simplePos x="0" y="0"/>
      <wp:positionH relativeFrom="margin"><wp:align>center</wp:align></wp:positionH>
      <wp:positionV relativeFrom="margin"><wp:align>top</wp:align></wp:positionV>
      <wp:extent cx="{width}" cy="{height}"/>
      <wp:effectExtent l="0" t="0" r="0" b="0"/>
      <wp:wrapTopAndBottom/>
      <wp:docPr id="{5000 + n}" name="{name}"/>
      <wp:cNvGraphicFramePr/>
      <a:graphic>
        <a:graphicData uri="{WPS}">
          <wps:wsp>
            <wps:cNvSpPr txBox="1"/>
            <wps:spPr>
              <a:xfrm><a:off x="0" y="0"/><a:ext cx="{width}" cy="{height}"/></a:xfrm>
              <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
              <a:noFill/>
              <a:ln><a:noFill/></a:ln>
            </wps:spPr>
            <wps:txbx><w:txbxContent/></wps:txbx>
            <wps:bodyPr rot="0" vert="horz" wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" anchor="t" anchorCtr="0">
              <a:spAutoFit/>
            </wps:bodyPr>
          </wps:wsp>
        </a:graphicData>
      </a:graphic>
    </wp:anchor>
  </w:drawing>
</w:r>"""
    run = etree.fromstring(xml)
    content = next(run.iter(W_NS + "txbxContent"))
    for el in elements:
        content.append(el)  # moves the elements out of the body
    if elements[-1].tag == W_NS + "tbl":  # a text box must end with a paragraph
        content.append(etree.SubElement(content, W_NS + "p"))
    ppr = anchor_el.find(W_NS + "pPr")
    if ppr is not None:
        ppr.addnext(run)
    else:
        anchor_el.insert(0, run)


# child order required by the OOXML schema (Word may refuse or "repair" files that violate it)
ORDER = {
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs "
           "suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd "
           "snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment "
           "textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
    "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid "
           "vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em "
           "lang eastAsianLayout specVanish oMath",
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing tblInd "
             "tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription",
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark",
    "sectPr": "headerReference footerReference footnotePr endnotePr type pgSz pgMar paperSrc pgBorders lnNumType "
              "pgNumType cols formProt vAlign noEndnote titlePg textDirection bidi rtlGutter docGrid printerSettings",
}
W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def fix_order(root):
    for tag, order in ORDER.items():
        rank = {name: k for k, name in enumerate(order.split())}
        for el in root.iter(W_NS + tag):
            kids = list(el)
            if not all(k.tag.startswith(W_NS) for k in kids):
                continue
            srt = sorted(kids, key=lambda k: rank.get(k.tag[len(W_NS):], len(rank)))
            if srt != kids:
                for k in kids:
                    el.remove(k)
                for k in srt:
                    el.append(k)


def drop_cap(p):
    """IEEE-style two-line drop cap for the first letter of the paragraph."""
    ts = list(p._p.iter(qn("w:t")))
    first = ts[0]
    letter, first.text = first.text[0], first.text[1:]
    # capitalise the rest of the first word, as \IEEEPARstart does
    word = re.match(r"\w*", first.text).group(0)
    first.text = word.upper() + first.text[len(word):]
    first.set(qn("xml:space"), "preserve")
    dp = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    fp = OxmlElement("w:framePr")
    for k, v in (("dropCap", "drop"), ("lines", "2"), ("wrap", "around"), ("vAnchor", "text"), ("hAnchor", "text")):
        fp.set(qn("w:" + k), v)
    ppr.append(fp)
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:line"), "455")
    sp.set(qn("w:lineRule"), "exact")
    ppr.append(sp)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:firstLine"), "0")
    ppr.append(ind)
    dp.append(ppr)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    pos = OxmlElement("w:position")
    pos.set(qn("w:val"), "-3")
    rpr.append(pos)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "56")
    rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = letter
    r.append(t)
    dp.append(r)
    p._p.addprevious(dp)
    p.paragraph_format.first_line_indent = Inches(0)


def main():
    build, out = Path(sys.argv[1]), Path(sys.argv[2])
    work = Path(tempfile.mkdtemp(prefix="fg_docx_"))
    try:
        figs, header, wide_tables = preprocess(build, work)
        raw = work / "raw.docx"
        subprocess.run(["pandoc", str(work / "pre.tex"), "-f", "latex", "-t", "docx", "-o", str(raw)], check=True,
                       cwd=work)
        postprocess(raw, out, figs, header, wide_tables)
        print(f"[docx] {out}")
    finally:
        if not os.environ.get("KEEP_WORK"):
            shutil.rmtree(work, ignore_errors=True)
        else:
            print("work dir:", work)


if __name__ == "__main__":
    main()
